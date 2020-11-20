#!/usr/bin/env python
#_*_coding:utf-8_*_

import docx
from docx import Document

docx_file = r"C:\Users\Administrator\Desktop\egatway.docx"
document = Document(docx_file)

tables = document.tables
print("该文档总共{tnum}个表格".format(tnum=len(tables)))

def get_params_data():
    params_data = {}

    ## 161 是参数前最后一个
    father_id = 0

    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        if tables.index(table_obj)< 162 or tables.index(table_obj)> 223:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))

            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)


                is_father = False
                if x[1]:
                    father_id = x[3]
                    is_father = True

                if not is_father:
                    # new_item["father_id"] = father_id
                    new_item.append(father_id)
                else:
                    new_item.append(None)
                params_data["%s^%s%s" % (x[3], x[4], "^%s" % x[5] if x[5] else "")] = new_item
            # print(x)
            # print(dir(table_obj))
            # print(table_obj.cell(i, 0).text)
            # print(table_obj.cell(i, 1).text)
            # print(table_obj.cell(i, 2).text)
            # print(table_obj.cell(i, 3).text)
            # print(table_obj.cell(i, 4).text)
            # print(table_obj.cell(i, 5).text)
            # print(table_obj.cell(i, 6).text)
        # break
    print(params_data)
    print(len(params_data))


def get_wave_data():
    params_data = {}

    ## 161 是参数前最后一个
    father_id = 0

    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        # if tables.index(table_obj) < 224 or tables.index(table_obj) > 227:
        if tables.index(table_obj) < 224 or tables.index(table_obj) > 228:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))

            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)

                params_data["%s^%s%s" % (x[1], x[2], "^%s" % x[3] if x[3] else "")] = new_item

    print(params_data)
    print(len(params_data))

def get_custom_wave_data():
    params_data = {}

    ## 161 是参数前最后一个
    father_id = 0

    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        if tables.index(table_obj) < 228 or tables.index(table_obj) > 228:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))

            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)

                params_data[x[1]] = new_item
    print(params_data)
    print(len(params_data))



def get_common_unit_data():
    params_data = {}

    ## 161 是参数前最后一个

    part_of = None
    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        # if tables.index(table_obj) < 229 or tables.index(table_obj) > 237:
        if tables.index(table_obj) < 229 or tables.index(table_obj) > 239:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))
            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)
                if not x[1].isdigit():
                    part_of = x[1]

                else:
                    new_item.append(part_of)
                    params_data["%s^%s%s" % (x[1], x[2], "^%s" % x[3] if x[3] else "")] = new_item
    print(params_data)
    print(len(params_data))

def get_custom_unit_data():
    params_data = {}

    ## 161 是参数前最后一个

    part_of = None
    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        if tables.index(table_obj) < 238 or tables.index(table_obj) > 239:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))
            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)
                is_part = False
                if not x[1].isdigit():
                    part_of = x[1]
                    is_part = True

                else:
                    new_item.append(part_of)
                    params_data[x[1]] = new_item
    print(params_data)
    print(len(params_data))


def get_body_data():
    params_data = {}

    ## 161 是参数前最后一个

    part_of = None
    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        if tables.index(table_obj) < 240 or tables.index(table_obj) > 244:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))
            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)
                if not x[1].isdigit():
                    part_of = x[1]

                else:
                    # new_item.append(part_of)
                    params_data["%s^%s%s" % (x[1], x[2], "^%s" % x[3] if x[3] else "")] = new_item
    print(params_data)
    print(len(params_data))


def get_enumerate_data():
    params_data = {}

    ## 161 是参数前最后一个

    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        if tables.index(table_obj) < 245 or tables.index(table_obj) > 259:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))
            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)
                if not x[1].isdigit():
                    part_of = x[1]

                else:
                    # new_item.append(part_of)
                    params_data["%s^%s%s" % (x[1], x[2], "^%s" % x[3] if x[3] else "")] = {new_item[0]:header[0], new_item[-1]:header[-1]}
    print(params_data)
    print(len(params_data))

def get_alarm_data():
    params_data = {}

    ## 161 是参数前最后一个

    for table_obj in tables:
        # header = table_obj[0]
        # if tables.index(table_obj) not in [223]:
        #     continue
        if tables.index(table_obj) < 260 or tables.index(table_obj) > 333:
            continue
        print("第{di_num}个表格".format(di_num=tables.index(table_obj)))
        header = []
        new_item = {}
        for i in range(0, len(table_obj.rows)):
            x = [table_obj.cell(i, j).text for j in range(len(table_obj.columns))]
            if i == 0:
                print("0行:%s" % (x))
                continue
            elif i == 1:
                header = x
                print("header:%s" % (x))
            else:

                # new_item = dict(zip(header,x))
                new_item = x
                print(new_item)
                if not x[1].isdigit():
                    part_of = x[1]

                else:
                    # new_item.append(part_of)
                    params_data["%s^%s%s" % (x[1], x[2], "^%s" % x[3] if x[3] else "")] = new_item
    print(params_data)
    print(len(params_data))

if __name__ == '__main__':
    # get_params_data()
    # get_wave_data()
    # get_custom_wave_data()
    # get_common_unit_data()
    # get_custom_unit_data()
    # get_body_data()
    # get_enumerate_data()
    get_alarm_data()